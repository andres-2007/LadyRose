"""HU 26: exportar la lista de inscritos a Excel."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

ENCABEZADOS = ["Nombre completo", "WhatsApp", "Correo", "Modalidad", "Estado de pago", "Fecha de inscripción"]


def libro_inscritos(cursos):
    libro = Workbook()
    libro.remove(libro.active)
    borde = Border(bottom=Side(style="thin", color="000000"))
    for curso in cursos:
        hoja = libro.create_sheet(curso.nombre[:31])
        hoja["A1"] = curso.nombre
        hoja["A1"].font = Font(bold=True, size=14)
        fecha = curso.fecha.strftime("%d/%m/%Y") if curso.fecha else "-"
        hoja["A2"] = (f"{curso.get_modalidad_display()} · {fecha} · {curso.horario} · "
                      f"Cupos: {curso.cupos_maximos} · Quedan: {curso.cupos_disponibles}")
        hoja.append([])
        hoja.append(ENCABEZADOS)
        for celda in hoja[4]:
            celda.font = Font(bold=True)
            celda.border = borde
        for i in curso.inscripciones.order_by("fecha"):
            hoja.append([i.nombre, i.whatsapp, i.correo, i.get_modalidad_display() or "-",
                         i.get_estado_pago_display(), i.fecha.strftime("%d/%m/%Y %H:%M")])
        for col, ancho in enumerate([30, 16, 30, 14, 16, 20], start=1):
            hoja.column_dimensions[get_column_letter(col)].width = ancho
        hoja.freeze_panes = "A5"
        for fila in hoja.iter_rows(min_row=5):
            for celda in fila:
                celda.alignment = Alignment(vertical="center")
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()
